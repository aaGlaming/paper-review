from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from backend.app.graph.review_graph import _parse_sections
from backend.app.importing.docx_text import docx_to_markdown


def _docx_bytes(document: Document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_headings_and_body_keep_document_order():
    document = Document()
    document.add_heading("摘要", level=1)
    document.add_paragraph("正文甲")
    document.add_heading("方法", level=2)
    document.add_paragraph("正文乙")

    result = docx_to_markdown(_docx_bytes(document))

    assert "# 摘要" in result and "## 方法" in result
    assert result.index("# 摘要") < result.index("正文甲") < result.index("## 方法") < result.index("正文乙")


def test_chinese_heading_style_name_is_recognized():
    document = Document()
    style = document.styles.add_style("标题 1", WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph("中文标题")
    paragraph.style = style
    assert "# 中文标题" in docx_to_markdown(_docx_bytes(document))


def test_plain_paragraphs_fall_back_to_unclassified_section():
    document = Document()
    document.add_paragraph("普通段落一")
    document.add_paragraph("普通段落二")

    result = docx_to_markdown(_docx_bytes(document))
    sections = _parse_sections(result)

    assert not result.startswith("#")
    assert len(sections) == 1
    assert sections[0]["section_type"] == "未分类"


def test_paragraph_table_paragraph_order_is_preserved():
    document = Document()
    document.add_paragraph("表格前段落")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格中内容"
    document.add_paragraph("表格后段落")

    result = docx_to_markdown(_docx_bytes(document))

    assert result.index("表格前段落") < result.index("表格中内容") < result.index("表格后段落")


def test_table_rows_columns_and_cell_newlines_are_preserved_as_markdown():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列一"
    table.cell(0, 1).text = "列|二"
    table.cell(0, 1).add_paragraph("换行内容")
    table.cell(1, 0).text = "甲"
    table.cell(1, 1).text = "乙"

    result = docx_to_markdown(_docx_bytes(document))
    table_lines = [line for line in result.splitlines() if line.startswith("|")]

    assert len(table_lines) == 3
    assert table_lines[0].count("|") == 3
    assert table_lines[1] == "| --- | --- |"
    assert "列 二 换行内容" in table_lines[0]
    assert "甲" in table_lines[2] and "乙" in table_lines[2]


def test_empty_table_is_not_dropped():
    document = Document()
    document.add_table(rows=0, cols=2)
    result = docx_to_markdown(_docx_bytes(document))
    assert result == "|  |\n| --- |"


@pytest.mark.parametrize("with_empty_paragraph", [False, True])
def test_empty_document_or_only_empty_paragraphs_raise(with_empty_paragraph):
    document = Document()
    if with_empty_paragraph:
        document.add_paragraph("")
    with pytest.raises(ValueError, match="docx"):
        docx_to_markdown(_docx_bytes(document))


def test_non_zip_bytes_raise_readable_docx_error():
    with pytest.raises(ValueError, match="docx"):
        docx_to_markdown(bytes(range(200)))


def test_outline_level_is_used_when_style_is_not_a_heading():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    paragraph = document.add_paragraph("結構標題")
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "1")
    paragraph._p.get_or_add_pPr().append(outline)
    assert "## 結構標題" in docx_to_markdown(_docx_bytes(document))
