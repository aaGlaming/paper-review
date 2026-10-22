from __future__ import annotations

import secrets
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool
from docx import Document

from backend.app.api import routes
from backend.app.api.app import create_app
from backend.app.auth import store as auth_store
from backend.app.auth.models import Permission, User, create_auth_schema
from backend.app.auth.security import hash_password
from backend.app.auth.token import issue_token
from backend.app.graph import review_graph
from backend.app.importing.docx_text import docx_to_markdown
from backend.app.report import models as report_models
from backend.app.report import store as report_store


@pytest.fixture()
def engine():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    create_auth_schema(engine)
    report_models.create_report_schema(engine)
    try:
        yield engine
    finally:
        engine.dispose()


@pytest.fixture()
def client(engine, tmp_path, monkeypatch):
    monkeypatch.setenv("AUTH_SECRET", secrets.token_urlsafe(48))
    monkeypatch.setattr(routes, "UPLOAD_ROOT", tmp_path / "uploads")
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    @contextmanager
    def _scope():
        session = factory()
        try:
            yield session
            session.commit()
        except Exception:
            session.rollback()
            raise
        finally:
            session.close()

    monkeypatch.setattr(auth_store, "session_scope_cm", _scope)

    class Graph:
        values = {}

        def get_state(self, _config):
            return SimpleNamespace(values=self.values)

    graph = Graph()

    def fake_create_review(paper_text, review_id, *, graph):
        graph.values = {"sections": review_graph._parse_sections(paper_text)}
        return review_id

    monkeypatch.setattr(review_graph, "create_review", fake_create_review)

    def _override_session():
        with _scope() as session:
            yield session

    app = create_app()
    app.dependency_overrides[routes.get_session] = _override_session
    app.dependency_overrides[routes.get_graph] = lambda: graph
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def _user(engine, username="student-upload"):
    with Session(engine) as session:
        row = auth_store.create_user(
            session,
            username=username,
            email=f"{username}@example.com",
            password_hash=hash_password("test-password-123"),
            role="student",
        )
        session.commit()
        return row.user_id


def _token(engine, user_id):
    with Session(engine) as session:
        row = session.get(User, user_id)
        return issue_token(row.user_id, row.role, row.password_hash)


def _headers(token):
    return {"Authorization": f"Bearer {token}"}


def _post(client, raw: bytes, filename="paper.md", token=None):
    headers = _headers(token) if token else {}
    return client.post(
        "/api/uploads",
        headers=headers,
        data={"contest_id": "contest-test"},
        files={"file": (filename, raw, "application/octet-stream")},
    )


def _counts(engine):
    with Session(engine) as session:
        return (
            session.scalar(select(func.count()).select_from(report_models.Review)) or 0,
            session.scalar(select(func.count()).select_from(report_models.ReviewSession)) or 0,
        )


def _files(root: Path):
    return list(root.rglob("*")) if root.exists() else []


def _docx_bytes(*, long_text="正文"):
    document = Document()
    document.add_heading("摘要", level=1)
    document.add_paragraph(long_text)
    document.add_heading("方法", level=2)
    document.add_paragraph("方法内容")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "甲"
    table.cell(0, 1).text = "乙"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_markdown_upload_creates_review_session_and_parses_atx(client, engine, monkeypatch):
    owner = _user(engine)
    source = "# 摘要\n结论文本\n\n## 方法\n方法内容\n"
    logged = []
    monkeypatch.setattr(routes, "log_event", lambda *args, **kwargs: logged.append((args, kwargs)))

    response = _post(client, source.encode("utf-8"), token=_token(engine, owner))

    assert response.status_code == 201
    body = response.json()
    assert body["paper_chars"] == len(source)
    assert body["source_sha256"]
    assert [section["title"] for section in body["sections"]] == ["摘要", "方法"]
    with Session(engine) as session:
        stored = session.get(report_models.ReviewSession, body["review_id"])
        assert stored.paper_text == source
        assert stored.user_id == owner
    target = Path(routes.UPLOAD_ROOT) / body["review_id"] / f"{body['source_sha256']}.md"
    assert target.read_bytes() == source.encode("utf-8")
    assert logged[0][0][:2] == ("importing", "paper_uploaded")
    assert set(logged[0][1]) == {"review_id", "bytes", "chars", "ext", "sha256", "source_format"}
    assert logged[0][1]["source_format"] == "md"


def test_utf8_bom_is_removed(client, engine):
    owner = _user(engine)
    response = _post(client, b"\xef\xbb\xbf# Title\nBody", token=_token(engine, owner))
    assert response.status_code == 201
    with Session(engine) as session:
        stored = session.get(report_models.ReviewSession, response.json()["review_id"])
        assert stored.paper_text.startswith("# Title")
        assert not stored.paper_text.startswith("\ufeff")


def test_gb18030_markdown_decodes_correctly(client, engine):
    owner = _user(engine)
    source = "# 摘要\n中文正文"
    response = _post(client, source.encode("gb18030"), token=_token(engine, owner))
    assert response.status_code == 201
    with Session(engine) as session:
        stored = session.get(report_models.ReviewSession, response.json()["review_id"])
        assert stored.paper_text == source


@pytest.mark.parametrize("raw", [b"", b" \r\n\t\n"])
def test_empty_or_whitespace_markdown_rejected_without_review_or_file(client, engine, raw):
    owner = _user(engine)
    before = _counts(engine)
    response = _post(client, raw, token=_token(engine, owner))
    assert response.status_code == 422
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_invalid_encoding_rejected_without_review_or_file(client, engine):
    owner = _user(engine)
    before = _counts(engine)
    response = _post(client, b"\xff\xfe\xff", token=_token(engine, owner))
    assert response.status_code == 422
    assert "文件编码不是有效 UTF-8/GB18030" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_byte_limit_is_injectable_and_rejected_before_writing(client, engine, monkeypatch):
    owner = _user(engine)
    monkeypatch.setattr(routes, "MAX_UPLOAD_BYTES", 5)
    before = _counts(engine)
    response = _post(client, b"123456", token=_token(engine, owner))
    assert response.status_code == 413
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_character_limit_is_injectable_and_rejected_before_writing(client, engine, monkeypatch):
    owner = _user(engine)
    monkeypatch.setattr(routes, "MAX_PAPER_CHARS", 3)
    before = _counts(engine)
    response = _post(client, "四字文本".encode("utf-8"), token=_token(engine, owner))
    assert response.status_code == 422
    assert "字符数" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


@pytest.mark.parametrize(
    ("suffix", "expected"),
    [
        (".txt", "仅支持 .md 与 .docx"),
        (".pdf", "系统不在线上解析 PDF"),
    ],
)
def test_non_markdown_extensions_are_rejected(client, engine, suffix, expected):
    owner = _user(engine)
    before = _counts(engine)
    response = _post(client, b"document", filename=f"paper{suffix}", token=_token(engine, owner))
    assert response.status_code == 415
    assert expected in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


@pytest.mark.anonymous
def test_missing_token_is_rejected_without_review_or_file(client, engine):
    before = _counts(engine)
    response = _post(client, b"# paper")
    assert response.status_code == 401
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_student_without_review_create_is_forbidden(client, engine):
    owner = _user(engine)
    with Session(engine) as session:
        permission = session.scalar(
            select(Permission).where(
                Permission.user_id == owner,
                Permission.permission == "review:create",
            )
        )
        session.delete(permission)
        session.commit()
    before = _counts(engine)
    response = _post(client, b"# paper", token=_token(engine, owner))
    assert response.status_code == 403
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_database_failure_removes_uploaded_file(client, engine, monkeypatch):
    owner = _user(engine)
    before = _counts(engine)
    logged = []
    monkeypatch.setattr(routes, "log_event", lambda *args, **kwargs: logged.append((args, kwargs)))

    def fail_save(*_args, **_kwargs):
        raise RuntimeError("simulated database failure")

    monkeypatch.setattr(report_store, "save_session", fail_save)
    response = _post(client, b"# paper\nbody", token=_token(engine, owner))

    assert response.status_code == 500
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []
    assert logged[-1][0][:2] == ("importing", "paper_upload_rollback")
    assert logged[-1][1]["file_removed"] is True


def test_docx_upload_creates_review_with_normalized_markdown(client, engine, monkeypatch):
    owner = _user(engine)
    raw = _docx_bytes()
    normalized = docx_to_markdown(raw)
    logged = []
    monkeypatch.setattr(routes, "log_event", lambda *args, **kwargs: logged.append((args, kwargs)))

    response = _post(client, raw, filename="paper.docx", token=_token(engine, owner))

    assert response.status_code == 201
    body = response.json()
    assert body["paper_chars"] == len(normalized)
    assert [section["title"] for section in body["sections"]] == ["摘要", "方法"]
    fetched = client.get(f"/api/reviews/{body['review_id']}", headers=_headers(_token(engine, owner)))
    assert fetched.status_code == 200
    assert len(fetched.json()["sections"]) == 2
    with Session(engine) as session:
        stored = session.get(report_models.ReviewSession, body["review_id"])
        assert stored.paper_text == normalized
    target = Path(routes.UPLOAD_ROOT) / body["review_id"] / f"{body['source_sha256']}.docx"
    assert target.read_bytes() == raw
    assert logged[-1][0][:2] == ("importing", "paper_uploaded")
    assert logged[-1][1]["source_format"] == "docx"
    assert "paper_text" not in logged[-1][1]


def test_docx_normalized_character_limit_rejects_before_review_or_file(client, engine, monkeypatch):
    owner = _user(engine)
    monkeypatch.setattr(routes, "MAX_PAPER_CHARS", 50)
    before = _counts(engine)

    response = _post(client, _docx_bytes(long_text="超" * 100), filename="paper.docx", token=_token(engine, owner))

    assert response.status_code == 422
    assert "字符数" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_corrupt_docx_rejects_without_review_or_file(client, engine):
    owner = _user(engine)
    before = _counts(engine)

    response = _post(client, bytes(range(200)), filename="paper.docx", token=_token(engine, owner))

    assert response.status_code == 422
    assert "docx" in response.text.lower()
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def _text_pdf_bytes():
    return Path("data/papers/cumcm-2024-C/C题.pdf").read_bytes()


def _scan_pdf_bytes():
    return Path("data/papers/cumcm-2024-A/A178.pdf").read_bytes()


def test_text_pdf_requires_offline_conversion_before_upload(client, engine):
    owner = _user(engine)
    raw = _text_pdf_bytes()
    before = _counts(engine)

    response = _post(client, raw, filename="paper.pdf", token=_token(engine, owner))

    assert response.status_code == 415
    assert "PDF 请先在外部转成文本或 Markdown 后再上传" in response.text
    assert "系统不在线上解析 PDF" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_pdf_upload_is_rejected_before_any_online_page_limit_check(client, engine):
    owner = _user(engine)
    before = _counts(engine)

    response = _post(client, _text_pdf_bytes(), filename="paper.pdf", token=_token(engine, owner))

    assert response.status_code == 415
    assert "系统不在线上解析 PDF" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_pdf_upload_is_rejected_before_any_online_page_count_check(client, engine):
    owner = _user(engine)
    before = _counts(engine)

    response = _post(client, _text_pdf_bytes(), filename="paper.pdf", token=_token(engine, owner))

    assert response.status_code == 415
    assert "系统不在线上解析 PDF" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_scanned_pdf_upload_is_rejected_before_online_text_layer_detection(client, engine):
    owner = _user(engine)
    before = _counts(engine)

    response = _post(client, _scan_pdf_bytes(), filename="paper.pdf", token=_token(engine, owner))

    assert response.status_code == 415
    assert "系统不在线上解析 PDF" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []


def test_corrupt_pdf_response_has_no_review_or_file(client, engine):
    owner = _user(engine)
    before = _counts(engine)

    response = _post(client, b"not a PDF\xff\x00", filename="paper.pdf", token=_token(engine, owner))

    assert response.status_code == 415
    assert "系统不在线上解析 PDF" in response.text
    assert _counts(engine) == before
    assert _files(routes.UPLOAD_ROOT) == []
